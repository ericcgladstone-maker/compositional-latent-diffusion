#!/usr/bin/env python3
"""Final targeted repairs to the formatted supplement.

Authorized changes only:
  1. S2.5 exponent -> true superscript 4 (formatting only; text unchanged)
  2-4. Three Table S1 "Range or levels" cells: replace TECHNICAL CHECK markers
       with the code-backed finding that low/high are nominal viewer metadata
  5. Remove the two decorative VML horizontal rules

NOT done: the S10.3 capacity label. Source says C=4, but the correct fix reaches
nine locations plus the Figure S2 artwork, which this pass may not touch. Left
verbatim and reported.
"""
import copy, shutil
from docx import Document
from docx.oxml.ns import qn

SRC = "submission/Diffusion_Before_Adoption_Supplement_Formatted.docx"
OUT = "submission/Diffusion_Before_Adoption_Supplement_Formatted_Final.docx"

RANGE_REPAIRS = {
    "capacity":
        "4 to 12 (nominal slider range; not enforced)",
    "change_agent_evangelism":
        "1.0 to 5.0 (nominal slider range; not enforced)",
    "replications":
        "1 to 100 (nominal slider range; not enforced)",
}

shutil.copy2(SRC, OUT)
doc = Document(OUT)
log = []

# ---- 1. superscript the exponent -----------------------------------------
target = None
for p in doc.paragraphs:
    if p.text.strip() == "1−0.854=0.478.":
        target = p; break
assert target is not None, "S2.5 exponent paragraph not found"
assert len(target.runs) == 1, f"expected 1 run, got {len(target.runs)}"
r0 = target.runs[0]
before, mid, after = "1−0.85", "4", "=0.478."
assert before + mid + after == r0.text

r0.text = before
new_runs = []
for txt, sup in ((mid, True), (after, False)):
    nr = copy.deepcopy(r0._r)
    for t in nr.findall(qn('w:t')):
        t.text = txt
        t.set(qn('xml:space'), 'preserve')
    rPr = nr.find(qn('w:rPr'))
    if rPr is None:
        rPr = nr.makeelement(qn('w:rPr'), {}); nr.insert(0, rPr)
    va = rPr.find(qn('w:vertAlign'))
    if sup:
        if va is None:
            va = rPr.makeelement(qn('w:vertAlign'), {}); rPr.append(va)
        va.set(qn('w:val'), 'superscript')
    elif va is not None:
        rPr.remove(va)
    r0._r.addnext(nr) if not new_runs else new_runs[-1].addnext(nr)
    new_runs.append(nr)
log.append(f"S2.5 exponent: split into 3 runs, '4' set vertAlign=superscript; "
           f"paragraph text unchanged ({target.text!r})")

# ---- 2-4. Table S1 range cells -------------------------------------------
t1 = doc.tables[24]
done = []
for row in t1.rows:
    name = row.cells[0].text.strip()
    if name in RANGE_REPAIRS and "TECHNICAL CHECK" in row.cells[3].text:
        cell = row.cells[3]
        old = cell.text
        paras = cell.paragraphs
        keep = paras[0]
        for extra in paras[1:]:
            extra._p.getparent().remove(extra._p)
        runs = keep.runs
        assert runs, f"no runs in {name} range cell"
        runs[0].text = RANGE_REPAIRS[name]
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
        done.append((name, old, RANGE_REPAIRS[name]))
assert len(done) == 3, f"expected 3 range repairs, applied {len(done)}"
for n, o, w in done:
    log.append(f"Table S1 [{n}] range: {o[:58]}... -> {w!r}")

# ---- 5. remove the two decorative VML rules ------------------------------
removed = 0
for p in list(doc.paragraphs):
    x = p._p.xml
    if '<w:pict' in x and 'v:rect' in x:
        assert not p.text.strip(), "rule paragraph unexpectedly carries text"
        for tag in ('<w:drawing', '<w:sectPr', 'bookmarkStart', 'commentRange',
                    '<w:hyperlink', '<w:fldChar'):
            assert tag not in x, f"rule paragraph also carries {tag}"
        p._p.getparent().remove(p._p)
        removed += 1
assert removed == 2, f"expected 2 VML rules, removed {removed}"
log.append(f"removed {removed} decorative VML horizontal rules "
           f"(text-empty paragraphs carrying only v:rect)")

doc.save(OUT)
print("wrote", OUT)
for l in log: print("  -", l)
