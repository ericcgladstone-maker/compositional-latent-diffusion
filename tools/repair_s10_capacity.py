#!/usr/bin/env python3
"""S10 advocate capacity label repair: C=5 -> C=4.

Source-backed: explore_changeagent.py and explore_changeagent_robust.py both set
REGIMES = {"permissive": 6, "constrained": 4}; stored output holds capacities 4
and 6 only. Labels only -- no value changes.

Leaves the legitimate C=5 references in S5/S6 (token-memory and joint/separate
designs genuinely use C=5,6) untouched.
"""
import hashlib, shutil, zipfile, os
from docx import Document

SRC = "submission/Diffusion_Before_Adoption_Supplement_Formatted_Final.docx"
OUT = "submission/Diffusion_Before_Adoption_Supplement_Final.docx"
FIG = "latent-diffusion-abm/results/figures-manuscript/figureS_advocates.png"

shutil.copy2(SRC, OUT)
doc = Document(OUT)
log = []

def swap_in_runs(para, old, new, label):
    """Replace within whichever run holds the substring; assert it is not split."""
    hits = 0
    for r in para.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            hits += r.text.count(new)
    assert hits, f"{label}: {old!r} not found intact in a single run"
    return hits

# ---- 1. S10.1 prose -------------------------------------------------------
n = swap_in_runs(doc.paragraphs[220], "At C=5,", "At C=4,", "S10.1 prose")
log.append(f"S10.1 prose (para 220): 'At C=5,' -> 'At C=4,' ({n})")

# ---- 2. S10.2 prose -------------------------------------------------------
n = swap_in_runs(doc.paragraphs[223], "At C=5,", "At C=4,", "S10.2 prose")
log.append(f"S10.2 prose (para 223): 'At C=5,' -> 'At C=4,' ({n})")

# ---- 3. S10.3 marker removal ---------------------------------------------
mk = doc.paragraphs[226]
assert mk.text.startswith("[TECHNICAL CHECK REQUIRED"), "marker paragraph moved"
assert "topology output" in mk.text
mk._p.getparent().remove(mk._p)
log.append("S10.3 marker paragraph removed in full")

# ---- 4. table headers -----------------------------------------------------
def set_cell(cell, new, label):
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    runs = paras[0].runs
    assert runs, f"{label}: no runs"
    old = cell.text
    runs[0].text = new
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    log.append(f"{label}: {old!r} -> {new!r}")

T = doc.tables
set_cell(T[21].rows[0].cells[3], "C = 4 adoption (SD)", "S10.1 table header [T22 0,3]")
set_cell(T[21].rows[0].cells[4], "C = 4 carriers",      "S10.1 table header [T22 0,4]")
set_cell(T[22].rows[0].cells[2], "C = 4 adoption (SD)", "S10.2 table header [T23 0,2]")
set_cell(T[23].rows[0].cells[3], "C = 4, no seed",       "S10.3 table header [T24 0,3]")
set_cell(T[23].rows[0].cells[4], "C = 4, seeded 0.20",   "S10.3 table header [T24 0,4]")

# ---- 5. Table S3 advocate rows only ---------------------------------------
for ri, name in ((11, "Advocate seeding"), (12, "Advocate activity")):
    cell = T[26].rows[ri].cells[3]
    assert T[26].rows[ri].cells[0].text.strip() == name, f"row {ri} is not {name}"
    assert "C=5,6" in cell.text
    p0 = cell.paragraphs[0]
    n = swap_in_runs(p0, "C=5,6", "C=4,6", f"Table S3 {name}")
    log.append(f"Table S3 [{name}]: 'C=5,6' -> 'C=4,6'")

# ---- 6. Figure S2 caption --------------------------------------------------
cap = doc.paragraphs[256]  # index shifts by 1 after the marker removal
assert cap.text.startswith("Figure S2."), f"caption moved: {cap.text[:40]!r}"
n1 = swap_in_runs(cap, "lower retention capacity (C=5)",
                  "lower retention capacity (C=4)", "Fig S2 caption label")
n2 = swap_in_runs(cap, "raises it modestly at C=5",
                  "raises it modestly at C=4", "Fig S2 caption second ref")
log.append("Figure S2 caption: 'lower retention capacity (C=5)' -> '(C=4)'")
log.append("Figure S2 caption: 'raises it modestly at C=5' -> 'at C=4'")

doc.save(OUT)

# ---- 7. swap the Figure S2 image binary (image2.png) ----------------------
newbytes = open(FIG, "rb").read()
tmp = OUT + ".tmp"
zin = zipfile.ZipFile(OUT); zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == "word/media/image2.png":
        data = newbytes
    zout.writestr(item, data)
zin.close(); zout.close(); os.replace(tmp, OUT)
log.append(f"Figure S2 binary (word/media/image2.png) replaced, "
           f"md5 -> {hashlib.md5(newbytes).hexdigest()}")

print("wrote", OUT)
for l in log: print("  -", l)
