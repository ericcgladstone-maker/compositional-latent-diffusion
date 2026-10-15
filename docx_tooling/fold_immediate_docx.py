#!/usr/bin/env python3
"""Fold the 2026-10-06 immediate-manifestation check into the build docx.

Round-4 comment C18 asks how results change if a ready agent manifests at once.
Two separate answers, kept separate on purpose:

  1. Main model, distributional robustness. Exogenous trigger every tick
     (`trigger_time=1`), carrier diagnostics and the capacity sweep at 200
     replications. Not a paired same-history contrast: the main engine draws
     activation from the same random stream as transmission.
  2. Three-arm token comparison, analytic invariance. Token histories never depend
     on manifestation, activation has its own stream, and fixed 1000-tick runs end
     on a scheduled trigger, so endpoint adoption equals endpoint readiness under
     either timing. `verify_token_activation_invariance.py` confirms it per rep.

Every value is computed from the result CSVs through the shared table builders.

  python fold_immediate_docx.py            # preview
  python fold_immediate_docx.py apply
"""
import os, re, shutil, sys

import docx
from docx.shared import Pt

import fold_200rep_tables as FT
from fold_carriers_source_docx import set_borders, BANNED

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.environ.get("ABM_DOCX") or os.path.abspath(
    os.path.join(HERE, "..", "latent-adoption-build-review.docx"))
ANCHOR = "New specification details"
LOCATOR = "Immediate manifestation upon readiness"


def _b(label):
    # builders are curried: factory() returns a function of the header row
    return FT.BUILDERS[label][0]()(None)


def table_carriers():
    b = _b("Immediate manifestation, carriers")
    hdr = ["Activation timing", "Carrier, at least 1, at initialization", "Carrier, at least 1, peak",
           "Carrier, at least 2, at initialization", "Carrier, at least 2, peak",
           "Share from nonmanifest sender"]
    return hdr, [[k] + v for k, v in b.items()]


def table_capacity():
    b = _b("Immediate manifestation, capacity")
    hdr = ["Capacity", "Adoption, scheduled", "Adoption, endogenous", "Adoption, immediate manifestation"]
    return hdr, [[k] + v for k, v in b.items()]


def table_invariance():
    b = _b("Token activation invariance")
    hdr = ["Readiness rule", "Capacity", "Loss rate", "Return to flow, scheduled",
           "Return to flow, immediate", "Largest per-replication difference"]
    return hdr, [list(k) + v for k, v in b.items()]


def values():
    c = _b("Immediate manifestation, carriers")
    k = _b("Immediate manifestation, capacity")
    t = _b("Token activation invariance")
    imm = c["Immediate"]
    return {
        "g1_0": imm[0], "g1_pk": imm[1], "g2_0": imm[2], "g2_pk": imm[3],
        "share": imm[4].split(" ")[0],
        "share_s": c["Scheduled"][4].split(" ")[0], "share_e": c["Endogenous"][4].split(" ")[0],
        "caps": ", ".join(k[c][2] for c in ("4", "5", "6")) + f" and {k['7'][2]} at capacities 4 to 7",
        "c6_i": k["6"][2], "c6_e": k["6"][1], "c6_s": k["6"][0],
        "maxdiff": max(float(v[2]) for v in t.values()),
    }


def para():
    v = values()
    return (
        f"{LOCATOR}. The round-4 review asked how the results change if a ready agent "
        "manifests at once rather than waiting for an activation condition. The engine "
        "expresses this by firing the exogenous trigger every tick, after readiness is "
        "recomputed, so readiness becomes manifestation within the same tick. Under this "
        f"timing the carrier fraction runs from {v['g1_0']} at initialization to a peak of "
        f"{v['g1_pk']}, the stricter two-element measure runs from {v['g2_0']} to {v['g2_pk']}, "
        f"and {v['share']} of successful relevant transfers come from a sender not manifest "
        f"for the big thing served, against {v['share_s']} under the scheduled trigger and "
        f"{v['share_e']} under the endogenous one. The lower values are expected, because an "
        "agent holding a sufficient configuration now counts as an adopter at once, so a "
        "nonmanifest sender is necessarily one below readiness. Adoption still rises with "
        f"capacity, {v['caps']}, and at capacity 6 it is {v['c6_i']} against {v['c6_e']} under "
        f"an endogenous trigger and {v['c6_s']} under the scheduled one. The comparison across "
        "timings is distributional rather than paired. The main engine draws activation from "
        "the same random stream as transmission, so changing the timing changes the element "
        "histories as well. The three-arm token comparison is unaffected by construction. Its "
        "token histories never depend on manifestation, activation has its own random stream, "
        "and the fixed 1000-tick runs end on a scheduled trigger, so final adoption equals "
        "readiness at the last tick under either timing. A replication-by-replication check "
        "over every cell finds the largest difference between the two timings to be "
        f"{v['maxdiff']:g}, and the scheduled arm reproduces the canonical per-replication values "
        f"exactly. From sweeps {FT.SRC['cdi']}, {FT.SRC['imc']} and {FT.SRC['tai']}, 200 "
        "replications each.")


def main():
    apply = len(sys.argv) > 1 and sys.argv[1] == "apply"
    p = para()
    bad = []
    if "—" in p:
        bad.append("em dash")
    if ";" in p:
        bad.append("semicolon")
    for b in BANNED:
        if re.search(r"\b" + re.escape(b) + r"\b", p.lower()):
            bad.append(f"banned term {b!r}")
    print("prose governance:", "OK" if not bad else "FAILED " + "; ".join(bad))
    if bad:
        sys.exit(1)
    print(f"\n--- paragraph ({len(p.split())} words)\n{p}")
    tabs = (table_carriers(), table_capacity(), table_invariance())
    for h, b in tabs:
        print("\n   " + " | ".join(h))
        for row in b:
            print("   " + " | ".join(row))
    if not apply:
        print("\npreview only, nothing written.")
        return
    doc = docx.Document(DOCX)
    if any(q.text.startswith(LOCATOR) for q in doc.paragraphs):
        sys.exit("already folded")
    bak = DOCX + ".bak21"
    if not os.path.exists(bak):
        shutil.copy2(DOCX, bak)
        print(f"\nbackup -> {os.path.basename(bak)}")
    anchor = next((q._p for q in doc.paragraphs if q.text.strip().startswith(ANCHOR)), None)
    if anchor is None:
        sys.exit("anchor not found")

    def mk_para(text):
        pp = doc.add_paragraph()
        pf = pp.paragraph_format
        pf.space_before = Pt(12); pf.space_after = Pt(6)
        a, b = text.split(". ", 1)
        pp.add_run(a + ". ").font.bold = True
        pp.add_run(b)
        return pp._p

    def mk_table(spec_):
        h, b = spec_
        t = doc.add_table(rows=1, cols=len(h))
        set_borders(t)
        for j, x in enumerate(h):
            t.rows[0].cells[j].paragraphs[0].add_run(x).font.bold = True
        for row in b:
            c = t.add_row().cells
            for j, x in enumerate(row):
                c[j].paragraphs[0].add_run(x)
        return t._tbl

    for el in [mk_para(p)] + [mk_table(t) for t in tabs]:
        anchor.addprevious(el)
    doc.save(DOCX)
    print("folded: 1 paragraph, 3 tables")


if __name__ == "__main__":
    main()
