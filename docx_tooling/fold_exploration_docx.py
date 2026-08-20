#!/usr/bin/env python3
"""Fold the wedge-exploration findings into the build-and-results .docx, IN PLACE.

Additive only: preserves existing prose and comments. Two folds:
  - By stage: the observability-gap figure + a short note into the wedge stage.
  - By finding: robustness and three-sources notes (+ collapse figure) into the
    closing synthesis.
Figures come from the verified exploration run; numbers are quoted from its log.
"""

import os
import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.abspath(os.path.join(HERE, "..", "latent-adoption-build-review.docx"))
EXP = os.path.join(HERE, "results", "20260623-144801-wedge-explore")
FIG_GAP = os.path.join(EXP, "fig_observability_gap.png")
FIG_COLLAPSE = os.path.join(EXP, "fig_collapse.png")

GAP_NOTE = (
    "Closer look, the observability gap. Tracking the substrate held (big-thing "
    "elements present in the population) against manifest adoption over time shows "
    "where the two transmissibility modes diverge. Until the first trigger the "
    "modes are identical, with the substrate rising to about 0.60 while adoption "
    "stays at zero, the latent build-up the argument names. After the trigger they "
    "become mirror images. Under elements only the substrate keeps leading adoption "
    "to the end, a positive gap near 0.09. Under whole package adoption overtakes "
    "the substrate, a negative gap near -0.17, because the manifest state itself is "
    "transmitted and a node can adopt without holding the full element set. The sign "
    "of the gap, not its size, separates the two. These trajectories come from a "
    "follow-up run at the same operating point, sweep 20260623-144801, 30 replications.")

ROBUST_NOTE = (
    "Robustness across structural variation. The wedge gap, whole package adoption "
    "minus elements only adoption, stays positive across every structural setting "
    "tested, topology, mean degree, population size, seeding density, transmission "
    "probability, and assembly difficulty, ranging about +0.18 to +0.62. Whole "
    "package does not always saturate. Where assembly is hard both modes fall, yet "
    "the gap does not close. The separation is a property of the construction rather "
    "than of the default settings. From sweep 20260623-144801, 30 replications.")

THREE_NOTE = (
    "Three sources of the wedge. The gap has three contributing barriers, assembly "
    "of multiple elements into a set, activation of a held substrate by a trigger, "
    "and competition for capacity among element types. Each one alone sustains a "
    "gap, and the gap reaches zero only when all three are removed together. The "
    "wedge is therefore the joint signature of the model's core pillars rather than "
    "an artifact of any single parameter. From sweep 20260623-144801.")


def find_p(body, prefix):
    for p in body.findall(qn("w:p")):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        if txt.startswith(prefix):
            return p
    raise SystemExit(f"anchor not found: {prefix!r}")


def build_para(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(12)
    if ". " in text:
        lead, rest = text.split(". ", 1)
        r1 = p.add_run(lead + ". "); r1.font.bold = True
        p.add_run(rest)
    else:
        p.add_run(text)
    return p._p


def build_fig(doc, path, width_in):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width_in))
    return p._p


def insert_after(anchor, elements):
    e = anchor
    for el in elements:
        e.addnext(el); e = el


def main():
    for f in (DOCX, FIG_GAP, FIG_COLLAPSE):
        if not os.path.exists(f):
            raise SystemExit(f"missing {f}")
    doc = docx.Document(DOCX)
    sec = doc.sections[0]
    fig_w = min(5.8, (sec.page_width - sec.left_margin - sec.right_margin) / 914400.0 * 0.92)
    body = doc.element.body

    # backup
    bak = DOCX + ".bak2"
    if not os.path.exists(bak):
        import shutil
        shutil.copy2(DOCX, bak)
        print(f"backup -> {bak}")

    # Fold 1: wedge stage, after the churn note
    churn = find_p(body, "Note on churn")
    insert_after(churn, [build_para(doc, GAP_NOTE),
                         build_fig(doc, FIG_GAP, fig_w)])

    # Fold 2: synthesis, after the last note paragraph
    last = find_p(body, "The visible trigger receives causal credit")
    insert_after(last, [build_para(doc, ROBUST_NOTE),
                        build_para(doc, THREE_NOTE),
                        build_fig(doc, FIG_COLLAPSE, fig_w)])

    doc.save(DOCX)
    print("folded: wedge-stage observability-gap note + figure; "
          "synthesis robustness note + three-sources note + collapse figure")


if __name__ == "__main__":
    main()
