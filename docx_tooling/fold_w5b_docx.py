#!/usr/bin/env python3
"""Fold the W5b loss-ladder ecology result into the docx (after the W5 ecology
material, before the porousness subsection). Prose-rule compliant. Table built
directly from the source CSV (traceable to raw). Honest negative: element loss
does not recover the recruitment-retention trade-off.
"""
import csv, glob, os, shutil
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.abspath(os.path.join(HERE, "..", "latent-adoption-build-review.docx"))
D3 = sorted(glob.glob(os.path.join(HERE, "results", "*-comments3")))[-1]
S3 = os.path.basename(D3).split("-comments3")[0]
w5b = list(csv.DictReader(open(os.path.join(D3, "w5b_loss_ladder.csv"))))

def f2(x): return f"{float(x):.2f}"

def w5b_table():
    losses = sorted({float(r["loss_rate"]) for r in w5b})
    betas = sorted({float(r["beta_compression"]) for r in w5b})
    hdr = ["Loss rate", "Rival compression", "Alpha share at compression 0", "Alpha share at compression 1", "Best-response compression"]
    body = []
    for L in losses:
        for b in betas:
            cells = [r for r in w5b if float(r["loss_rate"]) == L and float(r["beta_compression"]) == b]
            lo = next(r for r in cells if float(r["alpha_compression"]) == 0.0)
            hi = next(r for r in cells if float(r["alpha_compression"]) == 1.0)
            best = max(cells, key=lambda r: float(r["alpha_share"]))
            body.append([f"{L:g}", f"{b:g}", f2(lo["alpha_share"]), f2(hi["alpha_share"]),
                         f"{float(best['alpha_compression']):g}"])
    return hdr, body

PARA = ("Recruitment-retention under element loss. The competitive ecology was re-run across "
        "a ladder of element-loss rates, from 0 to 0.02, to test whether coupling recruitment "
        "to active members recovers the recruitment-retention trade-off. With no loss the run "
        "reproduces the ecology result above. Adding loss lowers adoption overall but does not "
        "change the shape. At every loss rate and every rival level, captured share still rises "
        "with compression and the best response is full compression, with no interior optimum in "
        "any of the twelve cells. Element loss thins the shared element pool uniformly and does "
        "not route recruitment through members, so encapsulation still carries no offsetting "
        "recruitment cost. Reproducing the trade-off would require recruitment to depend on "
        "active members, for example transmission gated on membership rather than on a shared "
        f"element pool. From sweep {S3}, 30 replications per cell.")

def set_borders(tbl):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    tbl._tbl.tblPr.append(borders)

def main():
    if not os.path.exists(DOCX):
        raise SystemExit("missing docx")
    bak = DOCX + ".bak7"
    if not os.path.exists(bak):
        shutil.copy2(DOCX, bak); print("backup ->", bak)
    doc = docx.Document(DOCX)
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Porousness and element overlap"):
            anchor = p._p; break
    if anchor is None:
        raise SystemExit("anchor 'Porousness and element overlap' not found")

    # build paragraph
    pp = doc.add_paragraph(); pf = pp.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(6)
    lead, rest = PARA.split(". ", 1)
    pp.add_run(lead + ". ").font.bold = True; pp.add_run(rest)
    # build table
    hdr, body = w5b_table()
    t = doc.add_table(rows=1, cols=len(hdr)); set_borders(t)
    for j, h in enumerate(hdr):
        t.rows[0].cells[j].paragraphs[0].add_run(h).font.bold = True
    for row in body:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].paragraphs[0].add_run(str(v))

    # insert both before the porousness paragraph, in order
    anchor.addprevious(pp._p)
    anchor.addprevious(t._tbl)
    doc.save(DOCX)
    print(f"folded W5b: 1 paragraph, 1 table ({len(body)} rows)")

if __name__ == "__main__":
    main()
