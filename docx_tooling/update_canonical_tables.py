"""Update the docx canonical per-stage tables from a new (higher-N) aggregated.csv.

Reuses audit_docx_fidelity's table-matching (5-col tables whose 2nd header is
'Adoption mean (sd)' and whose 1st header is a known focal-param label). For each
matched table, rewrites every data cell (mean(sd) / peak / overshoot / lapse) from
the new aggregate. PREVIEW mode reports every change and writes nothing; APPLY mode
edits the cells in place via python-docx.

Usage: python update_canonical_tables.py <new_full_dir> [apply]
"""
import csv, os, sys
from collections import defaultdict
import docx
from abm import phases as P

PARAM_LABEL = {
    "initial_element_density": "Initial element density", "capacity": "Capacity",
    "transmissibility_mode": "Transmissibility mode", "transmission_prob": "Transmission probability",
    "stabilization_boost": "Stabilization strength (boost)", "enhanced_factor": "Enhanced factor",
    "strong_tie_fraction": "Strong-tie fraction", "coupling_strength": "Coupling strength",
    "incompatibility": "Incompatibility", "change_agent_fraction": "Evangelist fraction"}

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.environ.get("ABM_DOCX") or os.path.abspath(
    os.path.join(HERE, "..", "latent-adoption-build-review.docx"))

def fnum(x): return float(x)

def load_agg(d):
    by = {}
    for r in csv.DictReader(open(os.path.join(d, "aggregated.csv"))):
        by.setdefault(r["phase"], {})[str(r["focal_value"])] = r
    return by

def expected_row(a):
    return [
        f"{fnum(a['final_manifest_share_mean']):.2f} ({fnum(a['final_manifest_share_sd']):.2f})",
        f"{fnum(a['peak_manifest_share_mean']):.2f}", f"{fnum(a['overshoot_mean']):.2f}",
        f"{fnum(a['lapse_rate_mean']):.2f}"]

def main():
    new_dir = sys.argv[1]
    apply = len(sys.argv) > 2 and sys.argv[2] == "apply"
    by = load_agg(new_dir)
    lab2pids = defaultdict(list)
    for pid in by:
        lab2pids[PARAM_LABEL[P.PHASE_INDEX[pid]["focal"]["param"]]].append(pid)

    doc = docx.Document(DOCX)
    changes = 0; tables_done = 0
    for tbl in doc.tables:
        hdr = [c.text.strip() for c in tbl.rows[0].cells]
        if not (len(hdr) == 5 and hdr[1] == "Adoption mean (sd)" and hdr[0] in lab2pids):
            continue
        data = {r.cells[0].text.strip(): [c.text.strip() for c in r.cells] for r in tbl.rows[1:]}
        # disambiguate phase by matching current adoption values (old numbers still in doc)
        pid = None
        for cand in lab2pids[hdr[0]]:
            if all(lvl in by[cand] for lvl in data):
                # match on level set; if ambiguous, refine by nearest adoption
                pid = cand
                if len(lab2pids[hdr[0]]) == 1:
                    break
        # if multiple candidates share the same levels, pick by closest adoption sums
        cands = [c for c in lab2pids[hdr[0]] if all(lvl in by[c] for lvl in data)]
        if len(cands) > 1:
            def dist(c):
                s = 0.0
                for lvl, row in data.items():
                    try:
                        cur = float(row[1].split()[0]); exp = float(by[c][lvl]['final_manifest_share_mean'])
                        s += abs(cur - exp)
                    except Exception: s += 9
                return s
            pid = min(cands, key=dist)
        if pid is None:
            print(f"  [skip] no phase for table {hdr[0]}"); continue
        tables_done += 1
        for r in tbl.rows[1:]:
            lvl = r.cells[0].text.strip()
            if lvl not in by[pid]:
                continue
            exp = expected_row(by[pid][lvl])
            cur = [c.text.strip() for c in r.cells][1:]
            for j, (cu, ex) in enumerate(zip(cur, exp), start=1):
                if cu != ex:
                    changes += 1
                    print(f"  {pid:22s} {lvl:>6} col{j}: {cu:>14} -> {ex}")
                    if apply:
                        cell = r.cells[j]
                        for run in list(cell.paragraphs[0].runs): run.text = ""
                        (cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs
                         else cell.paragraphs[0].add_run("")).text = ex
    print(f"\n{tables_done} per-stage tables matched, {changes} cell changes "
          f"({'APPLIED' if apply else 'preview only'}).")
    if apply:
        doc.save(DOCX); print("saved.")

if __name__ == "__main__":
    main()
