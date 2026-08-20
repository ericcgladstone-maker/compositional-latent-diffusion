#!/usr/bin/env python3
"""Update the number-bearing prose in the master docx to the 200-replication runs.

Design rule: every NEW numeric value is computed here from a result CSV on disk.
Nothing numeric is typed as a literal, so the doc can only ever be given a value
that exists in a source file. Non-numeric replacements (sweep ids, replication
wording) are literals and are listed explicitly.

Every edit asserts its OLD text is present exactly once in the target paragraph.
A missed edit is reported loudly rather than skipped silently.

Usage:
  python fold_200rep_prose.py            # preview, writes nothing
  python fold_200rep_prose.py apply
"""
import csv, json, os, sys
import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(HERE, "results")
DOCX = os.environ.get("ABM_DOCX") or os.path.abspath(
    os.path.join(HERE, "..", "latent-adoption-build-review.docx"))

SRC = {
    "canon":   "20260730-093338-full",
    "wedge":   "20260730-093438-wedge-explore",
    "tipping": "20260730-093440-tipping",
    "confirm": "20260730-093442-confirm",
    "w1238":   "20260730-102229-comments",
    "w4567":   "20260730-102232-comments2",
    "w5b":     "20260730-093336-comments3",
    "w5c":     "20260730-merged-comments4",
}
for _k in list(SRC):
    SRC[_k] = os.environ.get("SRC_" + _k, SRC[_k])


def rows(key, name):
    with open(os.path.join(RESULTS, SRC[key], name)) as f:
        return list(csv.DictReader(f))


def f2(x):
    return f"{float(x):.2f}"


def g(x):
    return f"{float(x):g}"


def half_capacity(pairs):
    """Linear interpolation of the capacity at which adoption crosses 0.5."""
    pairs = sorted(pairs)
    for (c0, v0), (c1, v1) in zip(pairs, pairs[1:]):
        if v0 < 0.5 <= v1:
            return c0 + (0.5 - v0) * (c1 - c0) / (v1 - v0)
    return None


# --------------------------------------------------------------- computed values
def compute():
    v = {}
    meta = json.load(open(os.path.join(RESULTS, SRC["canon"], "metadata.json")))
    reps = int(meta["replications"])
    agg = rows("canon", "aggregated.csv")
    v["reps"] = str(reps)
    v["total_runs"] = f"{len(agg) * reps:,}"
    v["n_settings"] = str(len(agg))

    # --- per-stage prose numbers, from the same canon the stage tables use ---
    # These sentences sit directly above the per-stage tables, so if they are not
    # updated with the tables the document contradicts itself on the same page.
    C = {}
    for r in agg:
        C[(r["phase"], str(r["focal_value"]))] = r

    def cv(phase, level, col="final_manifest_share_mean"):
        return f2(C[(phase, level)][col])

    def extreme(phase, col, how=max):
        vals = [float(r[col]) for r in agg if r["phase"] == phase]
        return f"{how(vals):.2f}"

    v["v1_tight"] = cv("V1", "4")
    v["v1_loose"] = cv("V1", "10")
    v["v1_lapse_max"] = extreme("V1", "lapse_rate_mean")
    v["vw_wp_lapse"] = cv("V1-W", "whole_package", "lapse_rate_mean")
    v["vw_eo_lapse"] = cv("V1-W", "elements_only", "lapse_rate_mean")
    v["vw_eo_over"] = cv("V1-W", "elements_only", "overshoot_mean")
    v["slot_lo"] = cv("V2-slotcost", "4")
    v["slot_hi"] = cv("V2-slotcost", "10")
    v["stab_full"] = cv("V2-stabilization", "0.0")
    v["stab_none"] = cv("V2-stabilization", "1.0")
    v["stab_over"] = extreme("V2-stabilization", "overshoot_mean")
    v["enh_over"] = cv("V2-enhanced", "3.0", "overshoot_mean")
    v["se_lo"] = cv("V2-stab-enh", "1.0")
    v["se_hi"] = cv("V2-stab-enh", "3.0")
    v["se_lapse_lo"] = cv("V2-stab-enh", "1.0", "lapse_rate_mean")
    v["se_lapse_hi"] = cv("V2-stab-enh", "3.0", "lapse_rate_mean")
    v["tie_lo"] = cv("V2-tiestrength", "0.0")
    v["tie_hi"] = cv("V2-tiestrength", "0.3")
    v["tie_over"] = extreme("V2-tiestrength", "overshoot_mean")
    v["inc_off"] = cv("V2-incompatibility", "False")
    v["inc_on"] = cv("V2-incompatibility", "True")

    # wedge gap range
    gaps = [float(r["wedge_gap"]) for r in rows("wedge", "part_b_wedge_gap.csv")]
    v["gap_min"] = f"{min(gaps):+.2f}"
    v["gap_max"] = f"{max(gaps):+.2f}"

    # bistability at caps 4,5,6 (random eviction) + closure capacity
    bi = [r for r in rows("tipping", "part_b_bistability.csv")
          if r.get("eviction", "random") == "random"]
    bmap = {g(r["capacity"]): float(r["bistability_gap"]) for r in bi}
    v["bi4"], v["bi5"], v["bi6"] = (f"{bmap['4']:.2f}", f"{bmap['5']:.2f}",
                                    f"{bmap['6']:.2f}")

    # capacity transition 4..8 and loss sensitivity
    a1 = {g(r["capacity"]): float(r["adopt_mean"]) for r in rows("tipping", "part_a1_aggregated.csv")}
    v["cap4"], v["cap5"], v["cap6"] = f2(a1["4"]), f2(a1["5"]), f2(a1["6"])
    v["cap7"], v["cap8"] = f2(a1["7"]), f2(a1["8"])
    a2 = {g(r["loss_rate"]): float(r["adopt_mean"]) for r in rows("tipping", "part_a2_aggregated.csv")}
    v["loss0"], v["loss01"] = f2(a2["0"]), f2(a2["0.01"])

    # strict-assembly bistability range (slack == 0)
    slack0 = [float(r["bistability_gap"]) for r in rows("confirm", "part1_slack.csv")
              if int(r["slack"]) == 0]
    v["slack_lo"], v["slack_hi"] = f"{min(slack0):.2f}", f"{max(slack0):.2f}"
    v["lapses_all_zero"] = all(float(r["lapses_above_mean"]) == 0
                               for r in rows("confirm", "part1_slack.csv"))

    # scale confirmation at N=1000
    sc = {(g(r["n_agents"]), r["mode"]): r for r in rows("confirm", "part2_wedge_scale.csv")}
    v["n1000_eo"] = f2(sc[("1000", "elements_only")]["final_manifest_share_mean"])
    v["n1000_wp"] = f"{float(sc[('1000','whole_package')]['final_manifest_share_mean']):.3f}"
    v["n1000_eo_gap"] = f2(sc[("1000", "elements_only")]["final_obs_gap_mean"])
    v["n1000_wp_gap"] = f2(sc[("1000", "whole_package")]["final_obs_gap_mean"])

    # W1 half-adoption capacities and the k/capacity ratio
    w1 = rows("w1238", "w1_capacity_x_k.csv")
    halves, ratios = {}, []
    for k in ("2", "3", "4"):
        pts = [(float(r["capacity"]), float(r["final_manifest_share"]))
               for r in w1 if g(r["sufficiency_k"]) == k]
        h = half_capacity(pts)
        halves[k] = h
        ratios.append(int(k) / h)
    v["half2"], v["half3"], v["half4"] = (f"{halves['2']:.1f}", f"{halves['3']:.1f}",
                                          f"{halves['4']:.1f}")
    v["ratio"] = f"{sum(ratios)/len(ratios):.2f}"

    # W2 corners
    w2 = {(g(r["stabilization_boost"]), g(r["enhanced_factor"])): r
          for r in rows("w1238", "w2_stab_x_enh.csv")}
    v["w2_lo"] = f2(w2[("1", "1")]["final_manifest_share"])
    v["w2_hi"] = f2(w2[("0", "3")]["final_manifest_share"])

    # W3 carriers
    w3 = {r["mode"]: r for r in rows("w1238", "w3_carriers.csv")}
    v["carrier_eo"] = f2(w3["elements_only"]["carrier_per_bt"])
    v["pure_eo"] = f2(w3["elements_only"]["pure_carrier"])
    v["carrier_wp"] = f2(w3["whole_package"]["carrier_per_bt"])

    # W8 density endpoints
    w8 = {g(r["initial_element_density"]): r for r in rows("w1238", "w8_density_under_cap.csv")}
    v["dens_lo"] = f2(w8["0.05"]["final_manifest_share"])
    v["dens_hi"] = f2(w8["0.6"]["final_manifest_share"])

    # W7 decomposition
    w7 = {r["mode"]: float(r["final"]) for r in rows("w4567", "w7_cuegated.csv")}
    v["w7_eo"], v["w7_gated"], v["w7_wp"] = (f2(w7["elements_only"]),
                                             f2(w7["whole_package_gated"]),
                                             f2(w7["whole_package"]))
    v["w7_trans"] = f2(w7["whole_package_gated"] - w7["elements_only"])
    v["w7_cue"] = f2(w7["whole_package"] - w7["whole_package_gated"])

    # W6 channels
    w6 = {(r["kind"], g(r["value"])): r for r in rows("w4567", "w6_coupling_matrix.csv")
          if r["kind"] in ("off_diagonal", "diagonal")}
    v["w6_base"] = f2(w6[("off_diagonal", "0")]["final"])
    v["w6_suppress"] = f2(w6[("off_diagonal", "-1")]["final"])
    v["w6_facilitate"] = f2(w6[("off_diagonal", "1")]["final"])
    v["w6_diag_hi"] = f2(w6[("diagonal", "1")]["final"])

    # W5 ecology endpoints
    w5 = {(g(r["beta_compression"]), g(r["alpha_compression"])): r
          for r in rows("w4567", "w5_ecology.csv")}
    v["w5_weak_lo"] = f2(w5[("0.2", "0")]["alpha_share"])
    v["w5_weak_hi"] = f2(w5[("0.2", "1")]["alpha_share"])
    v["w5_strong_lo"] = f2(w5[("0.8", "0")]["alpha_share"])
    v["w5_strong_hi"] = f2(w5[("0.8", "1")]["alpha_share"])

    # W4 leakage endpoints
    w4 = {g(r["element_overlap"]): r for r in rows("w4567", "w4_overlap_leakage.csv")}
    v["w4_lo"] = f2(w4["0"]["insular_leakage_to_beta"])
    v["w4_hi"] = f2(w4["3"]["insular_leakage_to_beta"])

    # W5b: cell count and whether any interior optimum exists
    lad = {}
    for r in rows("w5b", "w5b_loss_ladder.csv"):
        lad.setdefault((g(r["loss_rate"]), g(r["beta_compression"])), []).append(
            (float(r["alpha_compression"]), float(r["alpha_share"])))
    v["w5b_cells"] = str(len(lad))
    v["w5b_interior"] = sum(1 for c in lad.values()
                            if 0.0 < max(c, key=lambda t: t[1])[0] < 1.0)

    # W5c: interior count, grid size, loss levels that are interior
    mg = {}
    for r in rows("w5c", "w5c_member_gated.csv"):
        mg.setdefault((float(r["loss_rate"]), g(r["beta_compression"])), []).append(
            (float(r["alpha_compression"]), float(r["alpha_share"])))
    interior_losses = sorted({k[0] for k, c in mg.items()
                              if 0.0 < max(c, key=lambda t: t[1])[0] < 1.0})
    v["w5c_cells"] = str(len(mg))
    v["w5c_interior"] = sum(1 for c in mg.values()
                            if 0.0 < max(c, key=lambda t: t[1])[0] < 1.0)
    v["w5c_interior_losses"] = interior_losses
    # the span of loss rates at which the optimum is interior; at 200 reps this
    # reaches the top of the ladder, which it did not at 30
    v["w5c_loss_lo"] = f"{min(interior_losses):g}" if interior_losses else "none"
    v["w5c_loss_hi"] = f"{max(interior_losses):g}" if interior_losses else "none"
    v["w5c_floor"] = f2(min(v_ for c in mg.values() for _, v_ in c if _ == 1.0)
                        if any(a == 1.0 for c in mg.values() for a, _ in c) else 0.0)
    return v


NUM = {"0": "zero", "1": "one", "2": "two", "3": "three", "4": "four", "5": "five",
       "6": "six", "7": "seven", "8": "eight", "9": "nine", "10": "ten",
       "11": "eleven", "12": "twelve", "13": "thirteen", "14": "fourteen",
       "15": "fifteen", "16": "sixteen", "17": "seventeen", "18": "eighteen"}


def cite(key):
    """How a sweep is named in the document's prose.

    The doc's established convention is the bare timestamp ("From sweep
    20260623-144801"), not the full directory name. Keeping that matters for
    more than consistency: the directory suffixes include "-wedge-explore", and
    "wedge" is a banned term in this document's style guide, so citing full
    directory names would reintroduce it into the prose. The merged W5c result
    is not a single timed run, so it is named in full; PROVENANCE.md and the
    directory's own LOG.md record the six shards it came from.
    """
    d = SRC[key]
    stamp = d[:15]
    is_timestamped = (len(d) > 15 and stamp[:8].isdigit() and stamp[8] == "-"
                      and stamp[9:15].isdigit())
    return stamp if is_timestamped else d


def edits(v):
    """(paragraph locator, [(old, new), ...]). Locator must be unique in the doc."""
    W, T, C, B, M = (cite("wedge"), cite("tipping"), cite("confirm"),
                     cite("w1238"), cite("w4567"))
    L5B, L5C = cite("w5b"), cite("w5c")
    E = []

    # NB: "The build is progressive" alone also matches the design-rationale
    # paragraph in Implementation notes; anchor on the provenance sentence.
    E.append(("Final adoption is the mean manifest share at the end of a run", [
        ("30 replications per level", f"{v['reps']} replications per level"),
        ("1,620 seeded runs", f"{v['total_runs']} seeded runs")]))

    # --- per-stage prose, kept in step with the per-stage tables -------------
    E.append(("Limited attention. Capacity relative to assembly requirements", [
        ("Adoption rises from 0.11 under the tightest cap to 1.00",
         f"Adoption rises from {v['v1_tight']} under the tightest cap to {v['v1_loose']}"),
        ("reaching 0.35 under tight capacity",
         f"reaching {v['v1_lapse_max']} under tight capacity")]))

    E.append(("Churn under whole-package and elements-only transmission", [
        ("a lapse rate of 0.57", f"a lapse rate of {v['vw_wp_lapse']}"),
        ("a lapse rate of 0.24", f"a lapse rate of {v['vw_eo_lapse']}"),
        ("overshoot of 0.13", f"overshoot of {v['vw_eo_over']}")]))

    E.append(("Add-on: differential slot cost", [
        ("Adoption rises from 0.32 at the tightest capacity to 0.99",
         f"Adoption rises from {v['slot_lo']} at the tightest capacity to {v['slot_hi']}")]))

    E.append(("Add-on: stabilization. Protecting adopted elements", [
        ("reduces overshoot to approximately 0.05",
         f"reduces overshoot to approximately {v['stab_over']}"),
        ("adoption falls from 0.87 under full protection to 0.27 with no protection",
         f"adoption falls from {v['stab_full']} under full protection to "
         f"{v['stab_none']} with no protection")]))

    E.append(("Add-on: enhanced transmission", [
        ("overshoot rises to 0.21", f"overshoot rises to {v['enh_over']}")]))

    E.append(("Add-on: stabilization with enhanced transmission", [
        ("raises adoption from 0.39 to 0.66",
         f"raises adoption from {v['se_lo']} to {v['se_hi']}"),
        ("the lapse rate rises from 0.40 to 0.56",
         f"the lapse rate rises from {v['se_lapse_lo']} to {v['se_lapse_hi']}")]))

    E.append(("Add-on: tie strength channels", [
        ("from 0.51 with no strong ties to 0.73",
         f"from {v['tie_lo']} with no strong ties to {v['tie_hi']}"),
        ("overshoot no higher than 0.02", f"overshoot no higher than {v['tie_over']}")]))

    E.append(("Add-on: content-specific incompatibility", [
        ("reduces adoption from 0.65 to 0.44",
         f"reduces adoption from {v['inc_off']} to {v['inc_on']}")]))

    E.append(("Observability gap. Tracking the substrate held", [
        ("sweep 20260623-144801, 30 replications", f"sweep {W}, {v['reps']} replications")]))

    E.append(("Robustness across structural variation", [
        ("about +0.18 to +0.62", f"about {v['gap_min']} to {v['gap_max']}"),
        ("From sweep 20260623-144801, 30 replications.",
         f"From sweep {W}, {v['reps']} replications.")]))

    E.append(("Three contributions to the adoption gap", [
        ("From sweep 20260623-144801.", f"From sweep {W}.")]))

    E.append(("Bistability under reversible adoption", [
        ("with values 0.39, 0.19, and 0.35",
         f"with values {v['bi4']}, {v['bi5']}, and {v['bi6']}"),
        ("From sweep 20260623-210859, 30 replications per cell.",
         f"From sweep {T}, {v['reps']} replications per cell.")]))

    E.append(("Source of the bistability", [
        ("with values from 0.38 to 0.78",
         f"with values from {v['slack_lo']} to {v['slack_hi']}"),
        ("From sweep 20260623-214226, 50 replications per cell.",
         f"From sweep {C}, {v['reps']} replications per cell.")]))

    E.append(("Capacity transition and loss sensitivity", [
        ("with mean adoption 0.13, 0.31, 0.65, 0.99, and 1.00 from capacity 4 to 8",
         f"with mean adoption {v['cap4']}, {v['cap5']}, {v['cap6']}, {v['cap7']}, "
         f"and {v['cap8']} from capacity 4 to 8"),
        ("from 0.65 at no loss to 0.27 at a loss rate of 0.01",
         f"from {v['loss0']} at no loss to {v['loss01']} at a loss rate of 0.01"),
        ("From sweep 20260623-210859.", f"From sweep {T}.")]))

    E.append(("Confirmation at larger populations", [
        ("elements-only adoption is 0.64", f"elements-only adoption is {v['n1000_eo']}"),
        ("whole-package adoption is 0.996", f"whole-package adoption is {v['n1000_wp']}"),
        ("positive for elements-only at 0.11",
         f"positive for elements-only at {v['n1000_eo_gap']}"),
        ("negative for whole-package at -0.17",
         f"negative for whole-package at {v['n1000_wp_gap']}"),
        ("From sweep 20260623-214226, 50 replications per cell.",
         f"From sweep {C}, {v['reps']} replications per cell.")]))

    E.append(("Capacity transition and the needed-components ratio", [
        ("near 3.7, 5.6, and 7.2", f"near {v['half2']}, {v['half3']}, and {v['half4']}"),
        ("ratio near 0.54", f"ratio near {v['ratio']}"),
        ("From sweep 20260728-220913.", f"From sweep {B}.")]))

    E.append(("Stabilization and enhanced transmission, crossed", [
        ("from 0.27 at no protection and no enhancement to 1.00",
         f"from {v['w2_lo']} at no protection and no enhancement to {v['w2_hi']}"),
        ("From sweep 20260728-220913.", f"From sweep {B}.")]))

    E.append(("Carriers as a measured quantity", [
        ("the carrier fraction is 0.35 per big thing",
         f"the carrier fraction is {v['carrier_eo']} per big thing"),
        ("and 0.08 of agents hold", f"and {v['pure_eo']} of agents hold"),
        ("the carrier fraction is 0.01", f"the carrier fraction is {v['carrier_wp']}"),
        ("From sweep 20260728-220913.", f"From sweep {B}.")]))

    E.append(("Element density under a binding capacity", [
        ("from 0.65 at density 0.05 to 0.65 at density 0.60",
         f"from {v['dens_lo']} at density 0.05 to {v['dens_hi']} at density 0.60"),
        ("From sweep 20260728-220913.", f"From sweep {B}.")]))

    E.append(("Cue-gated whole-package, isolating transmissibility", [
        ("Final adoption is 0.65 under elements-only, 0.89 under cue-gated "
         "whole-package, and 0.99 under whole-package.",
         f"Final adoption is {v['w7_eo']} under elements-only, {v['w7_gated']} under "
         f"cue-gated whole-package, and {v['w7_wp']} under whole-package."),
        ("a transmissibility component of 0.24",
         f"a transmissibility component of {v['w7_trans']}"),
        ("a cue-gating component of 0.10", f"a cue-gating component of {v['w7_cue']}"),
        ("From sweep 20260728-223644.", f"From sweep {M}.")]))

    E.append(("Coupling matrix, separating the two channels", [
        ("adoption falls from 0.65 at no coupling to 0.51 under strong cross-target "
         "suppression and rises to 0.72 under cross-target facilitation",
         f"adoption falls from {v['w6_base']} at no coupling to {v['w6_suppress']} under "
         f"strong cross-target suppression and rises to {v['w6_facilitate']} under "
         f"cross-target facilitation"),
        ("nearly flat, from 0.65 to 0.67",
         f"nearly flat, from {v['w6_base']} to {v['w6_diag_hi']}"),
        ("From sweep 20260728-223644.", f"From sweep {M}.")]))

    E.append(("Competitive ecology and the recruitment-retention question", [
        ("from 0.40 to 0.74 against a weakly encapsulating rival and from 0.23 to 0.61 "
         "against a strongly encapsulating rival",
         f"from {v['w5_weak_lo']} to {v['w5_weak_hi']} against a weakly encapsulating "
         f"rival and from {v['w5_strong_lo']} to {v['w5_strong_hi']} against a strongly "
         f"encapsulating rival"),
        ("From sweep 20260728-223644.", f"From sweep {M}.")]))

    E.append(("Recruitment-retention under element loss", [
        (f"no interior optimum in any of the twelve cells",
         f"no interior optimum in any of the {NUM[v['w5b_cells']]} cells"),
        ("From sweep 20260729-081101, 30 replications per cell.",
         f"From sweep {L5B}, {v['reps']} replications per cell.")]))

    E.append(("Recruitment-retention with member-gated recruitment", [
        ("interior across loss rates from 0.002 to 0.006",
         f"interior across loss rates from {v['w5c_loss_lo']} to {v['w5c_loss_hi']}"),
        ("interior in eleven of the eighteen loss-by-rival cells",
         f"interior in {NUM[str(v['w5c_interior'])]} of the "
         f"{NUM[v['w5c_cells']]} loss-by-rival cells"),
        ("From sweep 20260729-090619, 30 replications per cell.",
         f"From sweep {L5C}, {v['reps']} replications per cell.")]))

    E.append(("Porousness and element overlap", [
        ("from 0.68 at no overlap to 0.83 at overlap 3",
         f"from {v['w4_lo']} at no overlap to {v['w4_hi']} at overlap 3"),
        ("From sweep 20260728-223644.", f"From sweep {M}.")]))

    return E


def content_runs(p):
    return [r for r in p.runs if r.text]


def set_para(p, lead, body):
    runs = content_runs(p)
    if not runs:
        return False
    bold = [r for r in runs if r.bold]
    norm = [r for r in runs if not r.bold]
    for t in p._p.iter(qn("w:t")):
        t.text = ""
    if lead and bold:
        bold[0].text = lead
        if norm:
            norm[0].text = body
        elif body:
            p.add_run(body)
    else:
        runs[0].text = (lead + body) if lead else body
    return True


def main():
    apply = len(sys.argv) > 1 and sys.argv[1] == "apply"
    v = compute()
    print("computed from source CSVs:")
    for k in sorted(v):
        print(f"   {k:22s} {v[k]}")

    doc = docx.Document(DOCX)
    paras = doc.paragraphs
    todo = edits(v)
    applied = missed = 0

    for locator, subs in todo:
        hits = [p for p in paras if locator in p.text]
        if len(hits) != 1:
            print(f"\n!! locator matched {len(hits)} paragraphs: {locator!r}")
            missed += len(subs)
            continue
        p = hits[0]
        runs = content_runs(p)
        bold = [r for r in runs if r.bold]
        lead = "".join(r.text for r in bold)
        full = "".join(r.text for r in runs)
        body = full[len(lead):]
        new_body = body
        print(f"\n--- {locator[:56]}")
        for old, new in subs:
            if old == new:
                print(f"    [unchanged] {old[:60]}")
                continue
            if new_body.count(old) != 1:
                print(f"    !! OLD not found exactly once ({new_body.count(old)}x): {old[:70]!r}")
                missed += 1
                continue
            new_body = new_body.replace(old, new)
            applied += 1
            print(f"    {old[:66]}\n      -> {new[:66]}")
        if apply and new_body != body:
            set_para(p, lead, new_body)

    print(f"\n{applied} prose replacements ({'APPLIED' if apply else 'preview only'}), "
          f"{missed} not applied.")
    if missed:
        print("REFUSING TO SAVE: some edits did not match." if apply else
              "(preview) some edits did not match.")
        if apply:
            sys.exit(1)
    if apply:
        doc.save(DOCX)
        print("saved.")


if __name__ == "__main__":
    main()
