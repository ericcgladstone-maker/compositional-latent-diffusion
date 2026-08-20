#!/usr/bin/env python3
"""Fold W5c (member-gated recruitment ecology) into the docx, after the W5b
subsection. Positive result: the recruitment-retention trade-off (inverted-U)
reproduces once recruitment runs through active members. Table built directly
from the source CSV. Prose-rule compliant.
"""
import csv, glob, os, shutil
import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.abspath(os.path.join(HERE, "..", "latent-adoption-build-review.docx"))
D4 = sorted(glob.glob(os.path.join(HERE, "results", "*-comments4")))[-1]
S4 = os.path.basename(D4).split("-comments4")[0]
FIG = os.path.join(D4, "fig_w5c_inverted_u.png")
rows = list(csv.DictReader(open(os.path.join(D4, "w5c_member_gated.csv"))))

def f2(x): return f"{float(x):.2f}"

def table():
    losses = sorted({float(r["loss_rate"]) for r in rows})
    betas = sorted({float(r["beta_compression"]) for r in rows})
    hdr = ["Loss rate", "Rival compression", "Best-response compression", "Alpha share at optimum", "Interior optimum"]
    body = []
    for L in losses:
        for b in betas:
            cells = [r for r in rows if float(r["loss_rate"]) == L and float(r["beta_compression"]) == b]
            best = max(cells, key=lambda r: float(r["alpha_share"]))
            bc = float(best["alpha_compression"])
            body.append([f"{L:g}", f"{b:g}", f"{bc:g}", f2(best["alpha_share"]),
                         "yes" if 0.0 < bc < 1.0 else "no"])
    return hdr, body

PARA = ("Recruitment-retention with member-gated recruitment. A second variant routes "
        "recruitment through members. With member-gated transmission on, a held element is "
        "transmitted only by an agent who is an active member of a big thing that element "
        "belongs to, so a compressed member holds no loose elements and transmits nothing, and "
        "encapsulation starves recruitment. Because member-routed recruitment cannot start from "
        "no members, each big thing is seeded with ten founders, as an organization begins with "
        "founders, and the ambient element density is set to zero so all spread is member-routed. "
        "Under these conditions the recruitment-retention trade-off appears. With no loss the best "
        "response is minimal compression, because recruitment dominates and retention carries no "
        "cost. As the loss rate rises the optimum moves to the interior, an inverted-U, with the "
        "best-response compression interior across loss rates from 0.002 to 0.006 and interior in "
        "eleven of the eighteen loss-by-rival cells. Full compression collapses each big thing to "
        "its ten founders, the 0.10 floor, because locked-in members recruit no one. Adoption is "
        "low in absolute terms under member-gated recruitment with loss, so the result is the "
        "shape rather than the level. The trade-off reported by Shi, Dokshin, Genkin, and "
        "<author> reproduces in this model once recruitment depends on active members, the "
        f"structural condition the earlier ecology runs lacked. From sweep {S4}, 30 replications per cell.")

def set_borders(tbl):
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        b.append(el)
    tbl._tbl.tblPr.append(b)

def main():
    for f in (DOCX, FIG):
        if not os.path.exists(f): raise SystemExit(f"missing {f}")
    bak = DOCX + ".bak9"
    if not os.path.exists(bak):
        shutil.copy2(DOCX, bak); print("backup ->", bak)
    doc = docx.Document(DOCX)
    sec = doc.sections[0]
    fig_w = min(5.6, (sec.page_width - sec.left_margin - sec.right_margin) / 914400.0 * 0.9)
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Porousness and element overlap"):
            anchor = p._p; break
    if anchor is None:
        raise SystemExit("anchor not found")
    # paragraph
    pp = doc.add_paragraph(); pf = pp.paragraph_format; pf.space_before = Pt(12); pf.space_after = Pt(6)
    lead, rest = PARA.split(". ", 1)
    pp.add_run(lead + ". ").font.bold = True; pp.add_run(rest)
    # table
    hdr, body = table()
    t = doc.add_table(rows=1, cols=len(hdr)); set_borders(t)
    for j, h in enumerate(hdr): t.rows[0].cells[j].paragraphs[0].add_run(h).font.bold = True
    for row in body:
        c = t.add_row().cells
        for j, v in enumerate(row): c[j].paragraphs[0].add_run(str(v))
    # figure
    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run().add_picture(FIG, width=Inches(fig_w))
    # insert all three before the porousness paragraph, in order
    anchor.addprevious(pp._p); anchor.addprevious(t._tbl); anchor.addprevious(fp._p)
    doc.save(DOCX)
    print(f"folded W5c: 1 paragraph, 1 table ({len(body)} rows), 1 figure")

if __name__ == "__main__":
    main()
