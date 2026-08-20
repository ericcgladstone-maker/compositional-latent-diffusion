#!/usr/bin/env python3
"""Fold the seeded-source sensitivity into the docx.

The distinct-senders readiness rule needs a source label for tokens that exist
before any transmission, and the focal specification gives each seeded token its
own origin. That is a convention, not a fact about the model, so the question is
whether the three-arm result depends on it. Two alternatives were run: one
exogenous origin per agent, and no source at all until the material has passed
through a transmission.

  python fold_seedsource_docx.py            # preview
  python fold_seedsource_docx.py apply
"""
import os, re, shutil, statistics as st, sys

import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import fold_200rep_tables as FT

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.environ.get("ABM_DOCX") or os.path.abspath(
    os.path.join(HERE, "..", "latent-adoption-build-review.docx"))
ANCHOR = "New specification details"
SRC_SS, SRC_SNR = FT.SRC["ss"], FT.SRC["snr"]

BANNED = ["wedge", "clean", "core", "central", "genuine", "interesting", "natural",
          "mirror images", "what we are seeing", "brain breaker", "sanity baseline"]

LABEL = {"unique": "unique origin per seeded token",
         "shared": "one exogenous origin per agent",
         "social_only": "no source until socially transmitted"}


def table():
    b = FT.BUILDERS["Seeded-source sensitivity"][0]()(None)
    hdr = ["Seed source convention", "Capacity", "Loss rate", "Return to flow",
           "Standard error"]
    body = []
    for mode in ("unique", "shared", "social_only"):
        for cap in ("6", "5"):
            for loss in ("0", "0.01"):
                body.append([LABEL[mode], cap, loss] + b[(mode, cap, loss)])
    return hdr, body


def para():
    b = FT.BUILDERS["Seeded-source sensitivity"][0]()(None)
    spread = []
    for cap in ("6", "5"):
        for loss in ("0", "0.01"):
            vals = [float(b[(m, cap, loss)][0]) for m in
                    ("unique", "shared", "social_only")]
            spread.append(max(vals) - min(vals))
    worst = max(spread)
    return (
        "Sensitivity to the seeded-source convention. The source rule needs an origin "
        "for every token, including the ones an agent holds before anything has been "
        "transmitted, and the focal specification gives each seeded token its own "
        "origin so that the initialization does not favour one readiness rule over "
        "another. That is a convention rather than a property of the model, so two "
        "alternatives were run. Under the first, everything an agent starts with "
        "counts as one exogenous origin. Under the second, seeded tokens carry no "
        "source at all and contribute nothing to source diversity until the material "
        "has passed through an actual transmission. The convention changes only the "
        "source labels. Component histories are byte-identical across the three "
        "settings, and the compositional and count arms are unaffected by it, which "
        "is asserted rather than assumed, so only the source arm was re-run. The "
        "three settings do differ where they should. At initialization the focal "
        "convention leaves about two agents per run already satisfying the source "
        "rule while the other two leave none, and at the first trigger adoption under "
        "the source rule is 0.442, 0.435 and 0.360 respectively. By the end of a run "
        f"the difference has gone. Returns to flow agree to within {worst:.4f} across "
        "all four cells, so the seeded-source convention governs the early dynamics "
        "and not the quantity the comparison rests on. Seeded material is displaced "
        "long before adoption is measured. The ordering of the three readiness rules "
        "is unchanged under every convention, with distinct components remaining the "
        "only rule whose response to greater transmission is negative under no loss. "
        f"From sweeps {SRC_SNR} and {SRC_SS}, 200 paired replications per cell.")


SPEC_ADD = (" The token-memory harness also exposes the seeded-source convention as a "
            "setting, with seeded tokens carrying either their own origin, one origin "
            "per agent, or no source at all until transmitted, which is what makes the "
            "sensitivity above expressible without touching component histories.")


def set_borders(tbl):
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        b.append(el)
    tbl._tbl.tblPr.append(b)


def main():
    apply = len(sys.argv) > 1 and sys.argv[1] == "apply"
    p = para()
    h, rows = table()
    bad = []
    for i, txt in enumerate((p, SPEC_ADD), 1):
        if "—" in txt:
            bad.append(f"text {i}: em dash")
        if ";" in txt:
            bad.append(f"text {i}: semicolon")
        for w in BANNED:
            if re.search(r"\b" + re.escape(w) + r"\b", txt.lower()):
                bad.append(f"text {i}: banned term {w!r}")
    print(f"sources: {SRC_SNR} | {SRC_SS}")
    print("prose governance:", "OK" if not bad else "FAILED " + "; ".join(bad))
    if bad:
        sys.exit(1)
    print(f"\n--- paragraph ({len(p.split())} words)\n{p}")
    print(f"\n--- table: {len(rows)} rows\n   " + " | ".join(h))
    for r in rows[:6]:
        print("   " + " | ".join(r))
    print(f"   ... {len(rows) - 6} more rows")
    if not apply:
        print("\npreview only, nothing written.")
        return

    bak = DOCX + ".bak20"
    if not os.path.exists(bak):
        shutil.copy2(DOCX, bak)
        print(f"\nbackup -> {os.path.basename(bak)}")
    doc = docx.Document(DOCX)
    spec = False
    for q in doc.paragraphs:
        if q.text.strip().startswith(ANCHOR) and SPEC_ADD.strip() not in q.text:
            q.add_run(SPEC_ADD); spec = True
    print("specification details extended:", spec)

    anchor = None
    for q in doc.paragraphs:
        if q.text.strip().startswith(ANCHOR):
            anchor = q._p
            break
    if anchor is None:
        sys.exit("anchor not found")

    pp = doc.add_paragraph()
    pf = pp.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(6)
    a, b = p.split(". ", 1)
    pp.add_run(a + ". ").font.bold = True
    pp.add_run(b)

    t = doc.add_table(rows=1, cols=len(h))
    set_borders(t)
    for j, x in enumerate(h):
        t.rows[0].cells[j].paragraphs[0].add_run(x).font.bold = True
    for row in rows:
        c = t.add_row().cells
        for j, x in enumerate(row):
            c[j].paragraphs[0].add_run(x)

    anchor.addprevious(pp._p)
    anchor.addprevious(t._tbl)
    doc.save(DOCX)
    print("folded: 1 paragraph, 1 table")


if __name__ == "__main__":
    main()
