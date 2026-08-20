#!/usr/bin/env python3
"""Fold the exploration results (tipping/bistability + confirmation) into the
master docx, and mark all comments resolved. Prose-rule compliant: no em dashes,
no semicolons, no performative language, numbers taken from the source CSVs.

Adds four subsections after the synthesis, with two figures. Comment resolution
sets w15:done="1" in commentsExtended.xml (comments are kept, marked resolved).
"""

import os, shutil, zipfile
import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.abspath(os.path.join(HERE, "..", "latent-adoption-build-review.docx"))
FIG_BISTAB = os.path.join(HERE, "results", "20260623-210859-tipping", "fig_bistability.png")
FIG_ASSEMBLY = os.path.join(HERE, "results", "20260623-214226-confirm", "fig_assembly_difficulty.png")

PARAS = [
 ("Bistability under reversible adoption. "
  "At a fixed capacity, a population started from a fresh low-adoption condition and a "
  "population started fully adopted reach different stable adoption levels. The bistability "
  "difference is positive across capacities 4 to 6, with values 0.39, 0.19, and 0.35, and it "
  "closes at capacity 7 and above where a fresh population also saturates. A fully adopted "
  "population is sustained at capacities where a fresh population does not assemble adoption. "
  "From sweep 20260623-210859, 30 replications per cell."),
 ("FIG", FIG_BISTAB),
 ("Source of the bistability. "
  "The bistability difference follows from a contrast between building an assembled set and "
  "maintaining one. Varying the assembly requirement at fixed capacity, the difference is "
  "largest under strict all-of assembly, where sufficiency equals the requirement size, with "
  "values from 0.38 to 0.78 across the requirement sizes tested. It falls to zero or slightly "
  "below as the threshold loosens, where a fresh population also reaches high adoption. The "
  "fully adopted state is maintained with zero recorded lapses, sustained by elements held "
  "across the population and not by repeated re-adoption. The difference depends on reversible "
  "adoption: with reversibility disabled a fresh population accumulates manifestations and the "
  "difference closes. From sweep 20260623-214226, 50 replications per cell."),
 ("FIG", FIG_ASSEMBLY),
 ("Capacity transition and loss sensitivity. "
  "Across replications the capacity transition is traversed over three integer capacity steps, "
  "with mean adoption 0.13, 0.31, 0.65, 0.99, and 1.00 from capacity 4 to 8. At capacities 5 "
  "and 6 the per-run outcomes spread across the range rather than splitting between zero and "
  "full adoption. Adoption falls continuously as the loss rate rises, from 0.65 at no loss to "
  "0.27 at a loss rate of 0.01 and near zero by 0.03. From sweep 20260623-210859."),
 ("Confirmation at larger populations. "
  "The adoption gap between whole-package and elements-only transmission, the reversal in the "
  "sign of the observability gap, and the capacity bistability were re-run at populations of "
  "500 and 1000 agents. At 1000 agents elements-only adoption is 0.64 and whole-package "
  "adoption is 0.996, the observability gap is positive for elements-only at 0.11 and negative "
  "for whole-package at -0.17, and the capacity bistability difference matches the "
  "smaller-population result. From sweep 20260623-214226, 50 replications per cell."),
]


def build_para(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(12)
    if ". " in text:
        lead, rest = text.split(". ", 1)
        r = p.add_run(lead + ". "); r.font.bold = True
        p.add_run(rest)
    else:
        p.add_run(text)
    return p._p


def build_fig(doc, path, w):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(w))
    return p._p


def resolve_comments(path):
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path, "r")
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    n = 0
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/commentsExtended.xml":
            n = data.count(b'w15:done="0"')
            data = data.replace(b'w15:done="0"', b'w15:done="1"')
        zout.writestr(item, data)
    zin.close(); zout.close()
    os.replace(tmp, path)
    return n


def main():
    for f in (DOCX, FIG_BISTAB, FIG_ASSEMBLY):
        if not os.path.exists(f):
            raise SystemExit(f"missing {f}")
    bak = DOCX + ".bak4"
    if not os.path.exists(bak):
        shutil.copy2(DOCX, bak)
        print(f"backup -> {bak}")

    doc = docx.Document(DOCX)
    sec = doc.sections[0]
    fig_w = min(5.8, (sec.page_width - sec.left_margin - sec.right_margin) / 914400.0 * 0.92)
    body = doc.element.body

    # anchor: end of synthesis
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Three contributions to the adoption gap"):
            anchor = p._p
            break
    if anchor is None:
        raise SystemExit("anchor 'Three contributions to the adoption gap' not found")

    # build elements in order
    new_els = []
    for item in PARAS:
        if isinstance(item, tuple) and item[0] == "FIG":
            new_els.append(build_fig(doc, item[1], fig_w))
        else:
            new_els.append(build_para(doc, item))

    e = anchor
    for el in new_els:
        e.addnext(el); e = el

    doc.save(DOCX)
    print(f"folded {sum(1 for x in PARAS if not (isinstance(x, tuple) and x[0]=='FIG'))} paragraphs "
          f"+ {sum(1 for x in PARAS if isinstance(x, tuple) and x[0]=='FIG')} figures")

    n = resolve_comments(DOCX)
    print(f"resolved comments: set done=1 on {n} entries")


if __name__ == "__main__":
    main()
